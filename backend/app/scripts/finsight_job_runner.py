from __future__ import annotations

import argparse
import asyncio
import json
import os
import shutil
import threading
import sys
import tempfile
import traceback
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

_STATUS_LOCK = threading.Lock()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run a FinSight report job for TickFlow.")
    parser.add_argument("--payload", required=True, help="Path to the request.json payload file.")
    parser.add_argument("--status", required=True, help="Path to the mutable status.json file.")
    parser.add_argument("--runs-root", required=True, help="Trusted parent directory for report runs.")
    return parser.parse_args()


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def load_payload(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Report payload must be a JSON object.")
    return payload


def _atomic_write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        with temporary.open("x", encoding="utf-8", newline="\n") as handle:
            json.dump(value, handle, indent=2, ensure_ascii=False)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_status(path: Path, patch: dict[str, Any]) -> None:
    with _STATUS_LOCK:
        current: dict[str, Any] = {}
        if path.exists():
            try:
                loaded = json.loads(path.read_text(encoding="utf-8"))
                current = loaded if isinstance(loaded, dict) else {}
            except (OSError, json.JSONDecodeError):
                current = {}
        if current.get("status") == "cancelled" and patch.get("status") != "cancelled":
            return
        current.update(patch)
        current["updated_at"] = now_iso()
        _atomic_write_json(path, current)


def validate_job_paths(
    payload: dict[str, Any],
    *,
    payload_path: Path,
    status_path: Path,
    runs_root: Path,
) -> None:
    run_dir = payload_path.parent.resolve()
    if payload_path.name != "request.json" or status_path.name != "status.json":
        raise ValueError("Unexpected report control filename.")
    if status_path.parent.resolve() != run_dir or run_dir.parent != runs_root.resolve():
        raise ValueError("Report control paths are outside the configured runs directory.")
    if Path(payload.get("run_dir", "")).resolve() != run_dir:
        raise ValueError("Payload run_dir does not match its control directory.")
    output_root = Path(payload.get("output_root", "")).resolve()
    if output_root.parent != run_dir:
        raise ValueError("Payload output_root is outside the run directory.")
    base_config_path = Path(payload.get("base_config_path", "")).resolve()
    trusted_finsight_root = Path.cwd().resolve()
    if (
        base_config_path.name != "my_config.yaml"
        or not base_config_path.is_file()
        or base_config_path.parent != trusted_finsight_root
        or base_config_path.is_symlink()
    ):
        raise ValueError("FinSight base configuration is unavailable.")
    source_root = trusted_finsight_root / "src"
    if not source_root.is_dir() or source_root.resolve() != source_root:
        raise ValueError("FinSight source directory is unavailable or redirected.")
    for key in ("run_id", "symbol", "stock_code", "market", "market_label", "target_name"):
        if not isinstance(payload.get(key), str) or not payload[key].strip():
            raise ValueError(f"Payload field {key} is required.")
    if payload["run_id"] != run_dir.name:
        raise ValueError("Payload run_id does not match its control directory.")
    for key in ("collect_tasks", "analysis_tasks"):
        tasks = payload.get(key)
        if not isinstance(tasks, list) or not tasks or not all(isinstance(item, str) and item.strip() for item in tasks):
            raise ValueError(f"Payload field {key} must be a non-empty list of strings.")


class ReportCancelled(RuntimeError):
    pass


def raise_if_cancelled(path: Path) -> None:
    if not path.exists():
        return
    try:
        current = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return
    if current.get("status") == "cancelled":
        raise ReportCancelled("Report generation cancelled by user.")


def choose_reference_doc(root: Path, base_config: dict[str, Any]) -> str:
    # TickFlow owns the final report branding. Start from the neutral template so
    # the FinSight Yulan/RUC header artwork is never carried into an attachment.
    preferred = root / "src" / "template" / "report_template.docx"
    if preferred.exists():
        return "src/template/report_template.docx"
    value = base_config.get("reference_doc_path")
    return str(value) if value else "src/template/report_template.docx"


def _copy_typeface_from_style(run: Any, style: Any) -> None:
    """Use the report body's typeface in page furniture without changing its size."""
    from copy import deepcopy
    from docx.oxml.ns import qn

    source_r_pr = style.element.find(qn("w:rPr"))
    source_r_fonts = source_r_pr.find(qn("w:rFonts")) if source_r_pr is not None else None
    if source_r_fonts is None:
        if style.font.name:
            run.font.name = style.font.name
        return

    destination_r_pr = run._element.get_or_add_rPr()
    for existing in destination_r_pr.findall(qn("w:rFonts")):
        destination_r_pr.remove(existing)
    destination_r_pr.append(deepcopy(source_r_fonts))


def _set_header_bottom_border(paragraph: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph_properties = paragraph._p.get_or_add_pPr()
    for border in paragraph_properties.findall(qn("w:pBdr")):
        paragraph_properties.remove(border)
    paragraph_border = OxmlElement("w:pBdr")
    bottom_border = OxmlElement("w:bottom")
    bottom_border.set(qn("w:val"), "single")
    bottom_border.set(qn("w:sz"), "10")
    bottom_border.set(qn("w:space"), "4")
    bottom_border.set(qn("w:color"), "C00000")
    paragraph_border.append(bottom_border)
    paragraph_properties.append(paragraph_border)


def _clear_and_write_header(header: Any, title: str, body_style: Any, title_color: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    from docx.shared import Pt

    paragraphs = list(header.paragraphs)
    paragraph = paragraphs[0] if paragraphs else header.add_paragraph()
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    _copy_typeface_from_style(run, body_style)
    run.font.color.rgb = title_color or RGBColor(0, 0, 0)
    run.font.size = Pt(9)
    _set_header_bottom_border(paragraph)


def _clear_and_write_footer(footer: Any, body_style: Any) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    paragraphs = list(footer.paragraphs)
    paragraph = paragraphs[0] if paragraphs else footer.add_paragraph()
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    paragraph.clear()
    paragraph_properties = paragraph._p.get_or_add_pPr()
    for border in paragraph_properties.findall(qn("w:pBdr")):
        paragraph_properties.remove(border)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    run = paragraph.add_run()
    _copy_typeface_from_style(run, body_style)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def apply_tickflow_report_style(docx_path: Path, target_name: str) -> None:
    """Apply TickFlow branding while preserving FinSight's document typography."""
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(docx_path)
    # Keep all report sections, but prevent a title followed by a level-two
    # heading from being split across two pages by the reference template.
    if "Heading 2" in document.styles:
        document.styles["Heading 2"].paragraph_format.page_break_before = False

    # Pandoc creates a live Word TOC. Restrict it to the two primary heading
    # levels so the final few entries do not spill onto an almost blank page.
    for instruction in document.element.body.iter(qn("w:instrText")):
        if instruction.text and 'TOC \\o "1-3"' in instruction.text:
            instruction.text = instruction.text.replace('TOC \\o "1-3"', 'TOC \\o "1-2"')

    body_style = document.styles["Body Text"] if "Body Text" in document.styles else document.styles["Normal"]
    title_color = document.styles["Title"].font.color.rgb if "Title" in document.styles else None
    header_title = f"{target_name}-投研报告"
    for section in document.sections:
        page_borders = section._sectPr.find(qn("w:pgBorders"))
        if page_borders is not None:
            section._sectPr.remove(page_borders)
        _clear_and_write_header(section.header, header_title, body_style, title_color)
        _clear_and_write_header(section.first_page_header, header_title, body_style, title_color)
        _clear_and_write_header(section.even_page_header, header_title, body_style, title_color)
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            _clear_and_write_footer(footer, body_style)

    document.save(docx_path)


def render_styled_pdf(docx_path: Path, pdf_path: Path) -> bool:
    """Render the post-processed DOCX so the downloadable PDF matches Word."""
    from docx2pdf import convert

    # Word's COM automation can fail on non-ASCII paths. The report workspace
    # normally contains Chinese directory names, so stage both conversion inputs
    # in the Windows system temp directory and copy only the finished PDF back.
    staging_root = Path(os.environ.get("SystemRoot", r"C:\\Windows")) / "Temp"
    staging_dir: Path | None = None
    try:
        if pdf_path.exists():
            pdf_path.unlink()
        staging_dir = Path(tempfile.mkdtemp(prefix="tickflow-pdf-", dir=staging_root))
        staged_docx = staging_dir / "report.docx"
        staged_pdf = staging_dir / "report.pdf"
        shutil.copy2(docx_path, staged_docx)
        convert(str(staged_docx), str(staged_pdf))
        # A valid research report PDF is far larger than FinSight's previous
        # one-page placeholder. Reject it before exposing an attachment.
        if not staged_pdf.exists() or staged_pdf.stat().st_size < 2048:
            return False
        shutil.copy2(staged_pdf, pdf_path)
        return True
    except Exception:
        return False
    finally:
        if staging_dir is not None:
            shutil.rmtree(staging_dir, ignore_errors=True)


def latest_report_file(working_dir: Path, suffix: str) -> Path | None:
    candidates = [
        path for path in working_dir.rglob(f"*{suffix}")
        if "outline" not in path.name.lower() and path.is_file() and path.stat().st_size > 0
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda item: item.stat().st_mtime)


def copy_artifact(src: Path | None, dst: Path) -> bool:
    if src is None or not src.exists() or src.stat().st_size <= 0:
        return False
    temporary = dst.with_name(f".{dst.name}.{uuid.uuid4().hex}.tmp")
    try:
        shutil.copy2(src, temporary)
        if temporary.stat().st_size <= 0:
            return False
        os.replace(temporary, dst)
        return True
    finally:
        temporary.unlink(missing_ok=True)


def infer_model_names(config) -> tuple[str, str, str]:
    llm_list = config.config.get("llm_config_list", [])
    names = [str(item.get("model_name", "")).strip() for item in llm_list if item.get("model_name")]
    llm_name = names[0] if len(names) >= 1 else ""
    embedding_name = names[1] if len(names) >= 2 else llm_name
    vlm_name = names[2] if len(names) >= 3 else llm_name
    return llm_name, embedding_name, vlm_name


def build_task_label(payload: dict[str, Any]) -> str:
    return (
        f"Research target: {payload['target_name']} "
        f"(market: {payload.get('market', payload['market_label'])}, ticker: {payload['stock_code']})"
    )


class StatusHeartbeat:
    def __init__(
        self,
        *,
        status_path: Path,
        stage: str,
        progress: int,
        message: str,
        interval: float = 10.0,
    ) -> None:
        self.status_path = status_path
        self.stage = stage
        self.progress = progress
        self.message = message
        self.interval = interval
        self._stop = threading.Event()
        self._thread: threading.Thread | None = None
        self._started_at = datetime.now()

    def start(self) -> None:
        self._thread = threading.Thread(target=self._run, daemon=True)
        self._thread.start()

    def stop(self) -> None:
        self._stop.set()
        if self._thread:
            self._thread.join(timeout=1.0)

    def _run(self) -> None:
        while not self._stop.wait(self.interval):
            if self.status_path.exists():
                try:
                    current = json.loads(self.status_path.read_text(encoding="utf-8"))
                    if current.get("status") == "cancelled":
                        return
                except Exception:
                    pass
            elapsed = int((datetime.now() - self._started_at).total_seconds())
            mins, secs = divmod(elapsed, 60)
            write_status(
                self.status_path,
                {
                    "status": "running",
                    "stage": self.stage,
                    "progress": self.progress,
                    "message": f"{self.message} running {mins:02d}:{secs:02d}",
                },
            )


async def run_job(payload: dict[str, Any], status_path: Path) -> None:
    raise_if_cancelled(status_path)
    base_config_path = Path(payload["base_config_path"]).resolve()
    root = base_config_path.parent

    os.chdir(root)
    sys.path.insert(0, str(root))

    from dotenv import load_dotenv

    load_dotenv(root / ".env", override=True)

    from src.agents import DataAnalyzer, DataCollector, ReportGenerator
    from src.config import Config
    from src.memory import Memory
    from src.utils import setup_logger

    base_config = {"reference_doc_path": "src/template/report_template.docx"}
    if base_config_path.exists():
        try:
            import yaml  # type: ignore

            base_config = yaml.safe_load(base_config_path.read_text(encoding="utf-8")) or {}
        except Exception:
            base_config = {"reference_doc_path": "src/template/report_template.docx"}

    config = Config(
        config_file_path=str(base_config_path),
        config_dict={
            "output_dir": payload["output_root"],
            "target_name": payload["target_name"],
            "stock_code": payload["stock_code"],
            "market": payload.get("market", ""),
            "target_type": "financial_company",
            "reference_doc_path": choose_reference_doc(root, base_config),
            "custom_collect_tasks": payload["collect_tasks"],
            "custom_analysis_tasks": payload["analysis_tasks"],
        },
    )
    memory = Memory(config=config)
    logger = setup_logger(log_dir=os.path.join(config.working_dir, "logs"))
    llm_name, embedding_name, vlm_name = infer_model_names(config)
    task_label = build_task_label(payload)

    total_steps = len(payload["collect_tasks"]) + len(payload["analysis_tasks"]) + 1
    completed = 0

    write_status(
        status_path,
        {
            "id": payload["run_id"],
            "symbol": payload["symbol"],
            "stock_code": payload["stock_code"],
            "market": payload.get("market", ""),
            "market_label": payload["market_label"],
            "name": payload["target_name"],
            "collect_tasks": payload["collect_tasks"],
            "analysis_tasks": payload["analysis_tasks"],
            "status": "running",
            "stage": "collecting",
            "progress": 5,
            "message": "FinSight job started",
            "started_at": now_iso(),
            "error": "",
            "pdf_status": "pending",
            "pdf_error": "",
            "pid": os.getpid(),
        },
    )

    # Match FinSight's original scheduler: tasks in one priority tier run in
    # parallel (up to three at a time), while later tiers wait for completion.
    priority_specs: dict[int, list[dict[str, Any]]] = {1: [], 2: []}
    for collect_task in payload["collect_tasks"]:
        priority_specs[1].append(
            {
                "label": collect_task,
                "agent_class": DataCollector,
                "task_input": {
                    "input_data": {
                        "task": f"{task_label}\n\nData collection task:\n{collect_task}"
                    },
                    "echo": True,
                    "max_iterations": 20,
                    "resume": False,
                },
                "agent_kwargs": {"use_llm_name": llm_name},
            }
        )
    for analysis_task in payload["analysis_tasks"]:
        priority_specs[2].append(
            {
                "label": analysis_task,
                "agent_class": DataAnalyzer,
                "task_input": {
                    "input_data": {"task": task_label, "analysis_task": analysis_task},
                    "echo": True,
                    "max_iterations": 20,
                    "resume": False,
                    "enable_chart": True,
                },
                "agent_kwargs": {
                    "use_llm_name": llm_name,
                    "use_vlm_name": vlm_name,
                    "use_embedding_name": embedding_name,
                },
            }
        )

    agents_by_priority: dict[int, list[dict[str, Any]]] = {1: [], 2: []}
    for priority, specs in priority_specs.items():
        for spec in specs:
            agent = await memory.get_or_create_agent(
                agent_class=spec["agent_class"],
                task_input=spec["task_input"],
                resume=False,
                priority=priority,
                **spec["agent_kwargs"],
            )
            agents_by_priority[priority].append({**spec, "agent": agent})
    memory.save()

    async def run_priority_group(priority: int, stage: str, message: str) -> None:
        nonlocal completed
        group = agents_by_priority[priority]
        if not group:
            return
        write_status(
            status_path,
            {
                "status": "running",
                "stage": stage,
                "progress": 5 + int((completed / total_steps) * 75),
                "message": message,
            },
        )
        logger.info(
            f"Priority {priority} group started: count={len(group)}, max_concurrent=3"
        )
        heartbeat = StatusHeartbeat(
            status_path=status_path,
            stage=stage,
            progress=5 + int((completed / total_steps) * 75),
            message=message,
        )
        heartbeat.start()
        semaphore = asyncio.Semaphore(3)

        async def run_one(item: dict[str, Any]) -> Any:
            nonlocal completed
            async with semaphore:
                raise_if_cancelled(status_path)
                logger.info(f"{stage} task started: {item['label']}")
                result = await item["agent"].async_run(**item["task_input"])
                completed += 1
                write_status(
                    status_path,
                    {
                        "status": "running",
                        "stage": stage,
                        "progress": 5 + int((completed / total_steps) * 75),
                        "message": f"{message} ({completed}/{total_steps - 1} tasks done)",
                    },
                )
                logger.info(f"{stage} task finished: {item['label']}")
                return result

        try:
            results = await asyncio.gather(*(run_one(item) for item in group), return_exceptions=True)
        finally:
            heartbeat.stop()
        failures = [result for result in results if isinstance(result, Exception)]
        if failures:
            raise RuntimeError(
                f"{stage} priority group failed: {failures[0]}"
            ) from failures[0]
        logger.info(f"Priority {priority} group finished")

    await run_priority_group(
        1,
        "collecting",
        f"Collecting {len(payload['collect_tasks'])} data tasks in parallel",
    )
    await run_priority_group(
        2,
        "analyzing",
        f"Analyzing {len(payload['analysis_tasks'])} tasks in parallel",
    )

    write_status(
        status_path,
        {
            "status": "running",
            "stage": "rendering",
            "progress": 88,
            "message": "Generating final report",
        },
    )
    raise_if_cancelled(status_path)
    logger.info("Report generation started")
    heartbeat = StatusHeartbeat(
        status_path=status_path,
        stage="rendering",
        progress=95,
        message="Generating final report",
    )
    heartbeat.start()
    try:
        report_input = {
            "input_data": {
                "task": task_label,
                "task_type": "company",
                "stock_code": payload["stock_code"],
            },
            "echo": True,
            "max_iterations": 20,
            "resume": False,
            "enable_chart": True,
        }
        agent = await memory.get_or_create_agent(
            agent_class=ReportGenerator,
            task_input=report_input,
            resume=False,
            priority=3,
            use_llm_name=llm_name,
            use_embedding_name=embedding_name,
        )
        report = await agent.async_run(**report_input)
    finally:
        heartbeat.stop()
    completed += 1

    working_dir = Path(config.working_dir)
    run_dir = Path(payload["run_dir"])
    md_path = latest_report_file(working_dir, ".md")
    docx_path = latest_report_file(working_dir, ".docx")
    if docx_path is None or not docx_path.exists():
        raise RuntimeError("FinSight completed but report.docx was not generated.")

    write_status(
        status_path,
        {
            "status": "running",
            "stage": "formatting",
            "progress": 97,
            "message": "Applying TickFlow report style",
        },
    )
    raise_if_cancelled(status_path)
    heartbeat = StatusHeartbeat(
        status_path=status_path,
        stage="formatting",
        progress=97,
        message="Applying TickFlow report style",
    )
    heartbeat.start()
    try:
        apply_tickflow_report_style(docx_path, payload["target_name"])
        styled_pdf_path = working_dir / "_tickflow_styled_report.pdf"
        styled_pdf_ready = render_styled_pdf(docx_path, styled_pdf_path)
    finally:
        heartbeat.stop()

    copied_md = copy_artifact(md_path, run_dir / "report.md")
    copied_docx = copy_artifact(docx_path, run_dir / "report.docx")
    copied_pdf = (
        copy_artifact(styled_pdf_path, run_dir / "report.pdf")
        if styled_pdf_ready
        else False
    )

    report_title = ""
    if copied_docx:
        report_title = docx_path.stem if docx_path else ""
    elif getattr(report, "title", None):
        report_title = str(report.title)

    write_status(
        status_path,
        {
            "status": "succeeded",
            "stage": "finished",
            "progress": 100,
            "message": "Report completed",
            "finished_at": now_iso(),
            "report_title": report_title,
            "pdf_status": "ready" if copied_pdf else "unavailable",
            "pdf_error": "" if copied_pdf else (
                "PDF was not generated. docx2pdf requires Microsoft Word in an interactive Windows session."
            ),
            "artifacts": {
                "markdown": copied_md,
                "word": copied_docx,
                "pdf": copied_pdf,
            },
        },
    )


def main() -> int:
    args = parse_args()
    payload_path = Path(args.payload).resolve()
    status_path = Path(args.status).resolve()
    runs_root = Path(args.runs_root).resolve()

    if (
        payload_path.name != "request.json"
        or status_path.name != "status.json"
        or payload_path.parent != status_path.parent
        or payload_path.parent.parent != runs_root
    ):
        print("Refusing report control paths outside the configured runs directory.", file=sys.stderr)
        return 2

    try:
        payload = load_payload(payload_path)
        validate_job_paths(
            payload,
            payload_path=payload_path,
            status_path=status_path,
            runs_root=runs_root,
        )
    except Exception as exc:  # noqa: BLE001
        write_status(
            status_path,
            {
                "status": "failed",
                "stage": "bootstrap_failed",
                "progress": 100,
                "message": str(exc),
                "error": traceback.format_exc(),
                "finished_at": now_iso(),
                "pdf_status": "unavailable",
                "pdf_error": "Runner failed before report generation started.",
            },
        )
        traceback.print_exc()
        return 1

    try:
        asyncio.run(run_job(payload, status_path))
        return 0
    except ReportCancelled as exc:
        write_status(
            status_path,
            {
                "status": "cancelled",
                "stage": "cancelled",
                "message": str(exc),
                "finished_at": now_iso(),
                "pdf_status": "unavailable",
                "pdf_error": "Report generation was cancelled by the user.",
            },
        )
        return 130
    except Exception as exc:  # noqa: BLE001
        write_status(
            status_path,
            {
                "status": "failed",
                "stage": "failed",
                "progress": 100,
                "message": str(exc),
                "error": traceback.format_exc(),
                "finished_at": now_iso(),
                "pdf_status": "unavailable",
                "pdf_error": "Report generation failed before a PDF could be produced.",
            },
        )
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
